import asyncio
import os
import sys

# Ensure we can import from backend
sys.path.insert(0, os.path.abspath('.'))

from resume.parser import ResumeParser

async def test_parser():
    parser = ResumeParser()
    
    # Create a dummy PDF using ReportLab or just use a dummy text
    # Actually, docx is easier to create or we can just mock a docx.
    
    # Let's try docx since python-docx is in requirements
    try:
        from docx import Document
        doc = Document()
        doc.add_paragraph("This is a resume of John Doe.")
        doc.add_paragraph("Skills: Python, Django, React, SQL")
        doc.add_paragraph("Experience: Software Engineer at Google")
        doc.save("test_resume.docx")
        
        print("Parsing docx...")
        result = await parser.parse_async("test_resume.docx")
        print("Parse Result:", result)
    except Exception as e:
        print("Error:", e)

if __name__ == "__main__":
    asyncio.run(test_parser())
